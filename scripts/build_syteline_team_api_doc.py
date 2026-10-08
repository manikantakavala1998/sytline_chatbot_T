"""Build the shareable SyteLine integration meeting handout from Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "documentation" / "SYTELINE_TEAM_API_REQUEST.md"
OUTPUT = ROOT / "documentation" / "SyteLine_Chatbot_API_Request_For_Meeting.docx"


def set_font(style, size: float, bold: bool = False) -> None:
    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Aptos")
    rfonts.set(qn("w:hAnsi"), "Aptos")


def rich_text(paragraph, value: str) -> None:
    """Render the limited bold and code syntax used in the source document."""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", value)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def clean_heading(value: str) -> str:
    value = re.sub(r"^(\d+)\.\s+", r"\1 ", value)
    value = value.replace(" — please work through these in order", "")
    value = value.replace("Explicitly later, not required for the first live-data release", "Later capabilities outside the first release")
    value = value.replace("Reuse the user's existing login", "Reuse the existing user login")
    return value.replace("—", " ").replace(":", "").replace(",", "").replace("-", " ").replace("'", "")


def suppress_paragraph_border(paragraph_or_style) -> None:
    ppr = paragraph_or_style._p.get_or_add_pPr() if hasattr(paragraph_or_style, "_p") else paragraph_or_style.element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is None:
        border = OxmlElement("w:pBdr")
        ppr.append(border)
    bottom = border.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        border.append(bottom)
    bottom.set(qn("w:val"), "nil")


def cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def cell_margins(cell, top=90, start=105, bottom=90, end=105) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = qn(f"w:{side}")
        node = margins.find(tag)
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:color"), "D9D9D9")


def add_interface_table(document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=0, cols=3)
    table.autofit = False
    widths = [Inches(1.48), Inches(0.82), Inches(4.85)]
    for row_index, data in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, data)):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            cell.paragraphs[0].style = "Table Body"
            rich_text(cell.paragraphs[0], value)
            if row_index == 0:
                cell_shading(cell, "24476B")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                cell_shading(cell, "F3F6F9")
        for index, width in enumerate(widths):
            cells[index].width = width
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = document.styles
    set_font(styles["Normal"], 10.4)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    set_font(styles["Title"], 20, True)
    styles["Title"].paragraph_format.space_after = Pt(11)
    suppress_paragraph_border(styles["Title"])
    set_font(styles["Heading 1"], 13, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    set_font(styles["Heading 2"], 11.2, True)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    set_font(styles["List Bullet"], 10.2)
    styles["List Bullet"].paragraph_format.left_indent = Inches(0.24)
    styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.12)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)
    if "Table Body" not in styles:
        styles.add_style("Table Body", 1)
    set_font(styles["Table Body"], 8.6)
    styles["Table Body"].paragraph_format.space_after = Pt(0)
    styles["Table Body"].paragraph_format.line_spacing = 1.05

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            suppress_paragraph_border(paragraph)
            paragraph.add_run("SyteLine Chatbot API and WebClient Integration Request")
        elif line.startswith("## "):
            document.add_paragraph(clean_heading(line[3:]), style="Heading 1")
        elif line.startswith("### "):
            if re.match(r"### (7|11)\. ", line):
                document.add_page_break()
            document.add_paragraph(clean_heading(line[4:]), style="Heading 2")
        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = []
            for table_line in table_lines:
                fields = [part.strip() for part in table_line.strip("|").split("|")]
                if all(re.fullmatch(r"[-: ]+", field or " ") for field in fields):
                    continue
                rows.append(fields)
            add_interface_table(document, rows)
            continue
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            rich_text(paragraph, line[2:])
        elif re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.first_line_indent = Inches(-0.2)
            paragraph.paragraph_format.space_after = Pt(3)
            rich_text(paragraph, line)
        else:
            paragraph = document.add_paragraph()
            rich_text(paragraph, line)
        index += 1

    document.core_properties.title = "SyteLine Chatbot API and WebClient Integration Request"
    document.core_properties.subject = "Read only API and WebClient integration discussion"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
