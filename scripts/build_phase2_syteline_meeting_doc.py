"""Create a concise Phase 2 integration request for the SyteLine team."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documentation" / "Phase_2_SyteLine_Integration_Questions.docx"


def style_font(style, size, *, bold=False):
    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Aptos")
    rfonts.set(qn("w:hAnsi"), "Aptos")


def remove_title_border(style, paragraph):
    for ppr in (style.element.get_or_add_pPr(), paragraph._p.get_or_add_pPr()):
        border = ppr.find(qn("w:pBdr"))
        if border is None:
            border = OxmlElement("w:pBdr")
            ppr.append(border)
        bottom = border.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            border.append(bottom)
        bottom.set(qn("w:val"), "nil")


def add_labeled_paragraph(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.add_run(label).bold = True
    paragraph.add_run(text)
    return paragraph


def add_step(doc, number, heading, example, question):
    doc.add_paragraph(f"{number} {heading}", style="Heading 2")
    add_labeled_paragraph(doc, "Example  ", example)
    add_labeled_paragraph(doc, "Question for the SyteLine team  ", question)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    style_font(doc.styles["Normal"], 10.6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    style_font(doc.styles["Title"], 18, bold=True)
    doc.styles["Title"].paragraph_format.space_after = Pt(11)
    style_font(doc.styles["Heading 2"], 11.6, bold=True)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)
    doc.styles["Heading 2"].paragraph_format.keep_with_next = True

    title = doc.add_paragraph(style="Title")
    remove_title_border(doc.styles["Title"], title)
    title.add_run("Phase 2 SyteLine Chatbot Integration Questions")

    intro = doc.add_paragraph()
    intro.add_run("Purpose  ").bold = True
    intro.add_run(
        "We need the chatbot to recognize the user already signed in to SyteLine, follow that user's access rights, "
        "and understand the WebClient screen they are viewing. Please confirm the supported integration methods for "
        "our installation. Live customer, order, and invoice data APIs are outside Phase 2."
    )

    doc.add_paragraph(
        "In the examples below, Ravi is signed in to SyteLine and opens the chatbot. His identity and screen context "
        "must come from trusted SyteLine and WebClient information, not from values typed into a chat request."
    )

    add_step(
        doc,
        1,
        "Verify the existing login",
        "Ravi opens the chatbot. The backend confirms that his SyteLine session is active and belongs to him.",
        "How can our backend verify the existing login without a second login? What happens when the session expires or Ravi logs out?",
    )
    add_step(
        doc,
        2,
        "Identify the current user and site",
        "Ravi belongs to the Sales group and is working in Site A. His individual access may differ from the group default.",
        "How do we obtain Ravi's user ID, active configuration and site, groups, allowed sites, and any individual access exceptions?",
    )
    add_step(
        doc,
        3,
        "Check effective permissions",
        "Ravi can open a form but cannot see one sensitive field. The chatbot must respect that restriction.",
        "What supported method tells us whether Ravi may see a form, field, or record? Will SyteLine enforce the same restrictions for API requests made as Ravi?",
    )
    add_step(
        doc,
        4,
        "Receive the current screen context",
        "Ravi is viewing a form with order 456 selected and asks, 'What does this field mean?' The chatbot needs the form, field, and selected record to understand the question.",
        "How can WebClient securely provide the active form, focused field, selected record key, and changes to them?",
    )
    add_step(
        doc,
        5,
        "Keep identity and access current",
        "Ravi changes sites, selects another record, logs out, or loses access after an administrator changes his permissions.",
        "How do we receive these changes, and how quickly must the chatbot stop using outdated session, screen, site, or permission information?",
    )

    doc.add_paragraph("Requested outcome", style="Heading 2")
    doc.add_paragraph(
        "Please provide the recommended integration method, documentation or interface contract, a redacted working example, "
        "and test users who show both allowed and denied access. We will use these to replace the chatbot's current simulated "
        "identity, permission, and screen context."
    )

    doc.core_properties.title = "Phase 2 SyteLine Chatbot Integration Questions"
    doc.core_properties.subject = "Existing login permissions and WebClient context"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
