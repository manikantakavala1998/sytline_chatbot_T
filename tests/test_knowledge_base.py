"""Guard the pre-Phase-4 knowledge corpus and sample Q&A boundary."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app.rag.markdown_processor import process_all_markdown


ROOT = Path(__file__).resolve().parents[1] / "data" / "knowledge" / "prospect_to_cash"
CORE_MODULES = {
    "prospect.md", "lead.md", "opportunity.md", "estimate.md", "quotation.md",
    "customer.md", "customer_order.md", "customer_order_line.md", "pricing.md",
    "credit.md", "shipment.md", "invoice.md", "payment.md", "faq.md",
    "campaigns_and_forecasts.md", "order_and_billing_variations.md",
}


def test_every_prospect_to_cash_article_is_chunked_with_unique_ids():
    files = {path.name for path in ROOT.glob("*.md")}
    chunks = process_all_markdown(ROOT)

    assert CORE_MODULES <= files
    assert {chunk.source_file for chunk in chunks} == files
    assert len({chunk.chunk_id for chunk in chunks}) == len(chunks)
    assert all(chunk.full_context_path and chunk.text for chunk in chunks)


def test_knowledge_articles_contain_complete_text_without_external_links():
    texts = [path.read_text(encoding="utf-8") for path in ROOT.glob("*.md")]

    assert all("http://" not in text and "https://" not in text and "](" not in text for text in texts)
    assert all("## " in text and "**Section Summary:**" in text for text in texts)


def test_prospect_qa_has_at_least_150_distinct_sourced_questions():
    workbook = load_workbook(ROOT.parents[1] / "qa" / "prospect_to_cash" / "prospect.xlsx", read_only=True)
    rows = list(workbook["qa_master"].values)
    headers = {name: index for index, name in enumerate(rows[0])}
    questions = rows[1:]
    new_rows = [row for row in questions if str(row[headers["qa_id"]]).startswith("PRSP-")]
    ids = [row[headers["qa_id"]] for row in questions]
    wording = [str(row[headers["canonical_question"]]).strip().casefold() for row in questions]
    variations = list(workbook["question_variations"].values)[1:]

    assert len(questions) >= 150
    assert len(ids) == len(set(ids))
    assert len(wording) == len(set(wording))
    assert all(row[headers["module"]] == "prospect" for row in new_rows)
    assert all(row[headers["source_reference"]] == "prospect.md" for row in new_rows)
    assert all(row[headers["approval_status"]] == "APPROVED" for row in new_rows)
    assert all(row[headers["active"]] is True for row in new_rows)
    assert all(row[0] in set(ids) for row in variations)


def test_customer_qa_has_at_least_150_distinct_screen_level_questions():
    workbook = load_workbook(ROOT.parents[1] / "qa" / "prospect_to_cash" / "customer.xlsx", read_only=True)
    rows = list(workbook["qa_master"].values)
    headers = {name: index for index, name in enumerate(rows[0])}
    questions = rows[1:]
    new_rows = [row for row in questions if str(row[headers["qa_id"]]).startswith("CUST-")]
    ids = [row[headers["qa_id"]] for row in questions]
    wording = [str(row[headers["canonical_question"]]).strip().casefold() for row in questions]
    variations = list(workbook["question_variations"].values)[1:]
    article = (ROOT / "customer.md").read_text(encoding="utf-8")

    assert len(new_rows) >= 150
    assert len(ids) == len(set(ids))
    assert len(wording) == len(set(wording))
    assert all(row[headers["module"]] == "customer" for row in new_rows)
    assert all(row[headers["source_reference"]] == "customer.md" for row in new_rows)
    assert all(row[headers["approval_status"]] == "APPROVED" for row in new_rows)
    assert all(row[headers["active"]] is True for row in new_rows)
    assert all(row[headers["answer"]] in article.replace("\n", " ") for row in new_rows)
    assert len(variations) == 3 * len(questions)
    assert all(row[0] in set(ids) for row in variations)


def test_customer_order_qa_has_at_least_150_distinct_header_questions():
    workbook = load_workbook(ROOT.parents[1] / "qa" / "prospect_to_cash" / "customer_order.xlsx", read_only=True)
    rows = list(workbook["qa_master"].values)
    headers = {name: index for index, name in enumerate(rows[0])}
    questions = rows[1:]
    new_rows = [row for row in questions if str(row[headers["qa_id"]]).startswith("CORD-")]
    ids = [row[headers["qa_id"]] for row in questions]
    wording = [str(row[headers["canonical_question"]]).strip().casefold() for row in questions]
    variations = list(workbook["question_variations"].values)[1:]
    headings = {
        line.removeprefix("## ").strip()
        for line in (ROOT / "customer_order.md").read_text(encoding="utf-8").splitlines()
        if line.startswith("## ")
    }

    assert len(new_rows) >= 150
    assert len(ids) == len(set(ids))
    assert len(wording) == len(set(wording))
    assert all(row[headers["module"]] == "customer_order" for row in new_rows)
    assert all(row[headers["source_reference"]] == "customer_order.md" for row in new_rows)
    assert all(row[headers["source_section"]].split(" > ")[-1] in headings for row in new_rows)
    assert all(row[headers["approval_status"]] == "APPROVED" for row in new_rows)
    assert all(row[headers["active"]] is True for row in new_rows)
    assert len(variations) == 3 * len(questions)


def test_customer_order_line_qa_has_at_least_150_distinct_line_questions():
    workbook = load_workbook(ROOT.parents[1] / "qa" / "prospect_to_cash" / "customer_order_line.xlsx", read_only=True)
    rows = list(workbook["qa_master"].values)
    headers = {name: index for index, name in enumerate(rows[0])}
    questions = rows[1:]
    new_rows = [row for row in questions if str(row[headers["qa_id"]]).startswith("COLN-")]
    ids = [row[headers["qa_id"]] for row in questions]
    wording = [str(row[headers["canonical_question"]]).strip().casefold() for row in questions]
    variations = list(workbook["question_variations"].values)[1:]
    article = (ROOT / "customer_order_line.md").read_text(encoding="utf-8")
    headings = {line[3:].strip() for line in article.splitlines() if line.startswith("## ")}

    assert len(new_rows) >= 150
    assert len(ids) == len(set(ids))
    assert len(wording) == len(set(wording))
    assert all(row[headers["module"]] == "customer_order_line" for row in new_rows)
    assert all(row[headers["source_reference"]] == "customer_order_line.md" for row in new_rows)
    assert all(row[headers["source_section"]].split(" > ")[-1] in headings for row in new_rows)
    assert all(row[headers["answer"]] in article for row in new_rows)
    assert all(row[headers["approval_status"]] == "APPROVED" for row in questions)
    assert all(row[headers["active"]] is True for row in questions)
    assert len(variations) == 3 * len(questions)


def test_pricing_qa_has_at_least_150_distinct_manual_questions():
    workbook = load_workbook(ROOT.parents[1] / "qa" / "prospect_to_cash" / "pricing.xlsx", read_only=True)
    rows = list(workbook["qa_master"].values)
    headers = {name: index for index, name in enumerate(rows[0])}
    questions = rows[1:]
    new_rows = [row for row in questions if str(row[headers["qa_id"]]).startswith("PRIC-")]
    ids = [row[headers["qa_id"]] for row in questions]
    wording = [str(row[headers["canonical_question"]]).strip().casefold() for row in questions]
    variations = list(workbook["question_variations"].values)[1:]
    article = (ROOT / "pricing.md").read_text(encoding="utf-8")
    headings = {line[3:].strip() for line in article.splitlines() if line.startswith("## ")}

    assert len(new_rows) >= 150
    assert len(ids) == len(set(ids))
    assert len(wording) == len(set(wording))
    assert all(row[headers["module"]] == "pricing" for row in new_rows)
    assert all(row[headers["source_reference"]] == "pricing.md" for row in new_rows)
    assert all(row[headers["source_section"]].split(" > ")[-1] in headings for row in new_rows)
    assert all(row[headers["answer"]] in article for row in new_rows)
    assert all("User Guide 9.01.x" in row[headers["version"]] for row in new_rows)
    assert all(row[headers["approval_status"]] == "APPROVED" for row in questions)
    assert all(row[headers["active"]] is True for row in questions)
    assert len(variations) == 3 * len(questions)


def test_credit_qa_has_at_least_150_distinct_manual_questions():
    workbook = load_workbook(ROOT.parents[1] / "qa" / "prospect_to_cash" / "credit.xlsx", read_only=True)
    rows = list(workbook["qa_master"].values)
    headers = {name: index for index, name in enumerate(rows[0])}
    questions = rows[1:]
    new_rows = [row for row in questions if str(row[headers["qa_id"]]).startswith("CRED-")]
    ids = [row[headers["qa_id"]] for row in questions]
    wording = [str(row[headers["canonical_question"]]).strip().casefold() for row in questions]
    variations = list(workbook["question_variations"].values)[1:]
    article = (ROOT / "credit.md").read_text(encoding="utf-8")
    headings = {line[3:].strip() for line in article.splitlines() if line.startswith("## ")}

    assert len(new_rows) >= 150
    assert len(ids) == len(set(ids))
    assert len(wording) == len(set(wording))
    assert all(row[headers["module"]] == "credit" for row in new_rows)
    assert all(row[headers["source_reference"]] == "credit.md" for row in new_rows)
    assert all(row[headers["source_section"]].split(" > ")[-1] in headings for row in new_rows)
    assert all(row[headers["answer"]] in article for row in new_rows)
    assert all("User Guide 9.01.x" in row[headers["version"]] for row in new_rows)
    assert all(row[headers["approval_status"]] == "APPROVED" for row in questions)
    assert all(row[headers["approved_by"]] == "User-authorized automatic approval (SME review pending)" for row in questions)
    assert all(row[headers["active"]] is True for row in questions)
    assert len(variations) == 3 * len(questions)


def test_shipment_qa_has_at_least_150_distinct_manual_questions():
    workbook = load_workbook(ROOT.parents[1] / "qa" / "prospect_to_cash" / "shipment.xlsx", read_only=True)
    rows = list(workbook["qa_master"].values)
    headers = {name: index for index, name in enumerate(rows[0])}
    questions = rows[1:]
    new_rows = [row for row in questions if str(row[headers["qa_id"]]).startswith("SHIP-")]
    ids = [row[headers["qa_id"]] for row in questions]
    wording = [str(row[headers["canonical_question"]]).strip().casefold() for row in questions]
    variations = list(workbook["question_variations"].values)[1:]
    article = (ROOT / "shipment.md").read_text(encoding="utf-8")
    headings = {line[3:].strip() for line in article.splitlines() if line.startswith("## ")}

    assert len(new_rows) >= 150
    assert len(ids) == len(set(ids))
    assert len(wording) == len(set(wording))
    assert all(row[headers["module"]] == "shipment" for row in new_rows)
    assert all(row[headers["source_reference"]] == "shipment.md" for row in new_rows)
    assert all(row[headers["source_section"]].split(" > ")[-1] in headings for row in new_rows)
    assert all(row[headers["answer"]] in article for row in new_rows)
    assert all("User Guide 9.01.x" in row[headers["version"]] for row in new_rows)
    assert all(row[headers["approval_status"]] == "APPROVED" for row in questions)
    assert all(row[headers["approved_by"]] == "User-authorized automatic approval (SME review pending)" for row in questions)
    assert all(row[headers["active"]] is True for row in questions)
    assert len(variations) == 3 * len(questions)


@pytest.mark.parametrize(
    ("module", "prefix", "manual"),
    [
        ("invoice", "INV-", "Customer Service User Guide 9.01.x"),
        ("payment", "PAY-", "Financials User Guide 9.01.x"),
    ],
)
def test_invoice_and_payment_qa_are_distinct_sourced_and_user_activated(module, prefix, manual):
    qa_path = ROOT.parents[1] / "qa" / "prospect_to_cash" / f"{module}.xlsx"
    workbook = load_workbook(qa_path, read_only=True, data_only=True)
    rows = list(workbook["qa_master"].values)
    headers = {name: index for index, name in enumerate(rows[0])}
    questions = rows[1:]
    new_rows = [row for row in questions if str(row[headers["qa_id"]]).startswith(prefix)]
    ids = [row[headers["qa_id"]] for row in questions]
    wordings = [str(row[headers["canonical_question"]]).strip().casefold() for row in questions]
    variations = list(workbook["question_variations"].values)[1:]
    article = (ROOT / f"{module}.md").read_text(encoding="utf-8")
    headings = {line[3:].strip() for line in article.splitlines() if line.startswith("## ")}

    assert len(new_rows) >= 150
    assert len(ids) == len(set(ids))
    assert len(wordings) == len(set(wordings))
    assert all(row[headers["module"]] == module for row in questions)
    assert all(row[headers["source_reference"]] == f"{module}.md" for row in questions)
    assert all(row[headers["source_section"]].split(" > ")[-1] in headings for row in questions)
    assert all(row[headers["answer"]] in article for row in new_rows)
    assert all(manual in row[headers["version"]] for row in questions)
    assert all(row[headers["approval_status"]] == "APPROVED" for row in questions)
    assert all(
        row[headers["approved_by"]] == "User-authorized automatic approval (SME review pending)"
        for row in questions
    )
    assert all(row[headers["active"]] is True for row in questions)
    assert len(variations) == 3 * len(questions)
    assert all(row[0] in set(ids) for row in variations)
