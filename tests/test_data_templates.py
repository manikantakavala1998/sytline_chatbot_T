"""Data-team templates stay loadable: Markdown/Word/Excel templates, the checker and the converter."""

import shutil
from pathlib import Path

import openpyxl

from backend.app.qa.loader import _load_one_file
from backend.app.rag.markdown_processor import build_chunks
from scripts import convert_docx_to_md, validate_knowledge_data as check

TEMPLATES = Path(__file__).resolve().parents[1] / "documentation" / "data_templates"
MD_TEMPLATE = TEMPLATES / "SyteLine_Module_Knowledge_Template.md"
DOCX_TEMPLATE = TEMPLATES / "SyteLine_Module_Knowledge_Template.docx"
XLSX_TEMPLATE = TEMPLATES / "SyteLine_QA_Template.xlsx"
EXAMPLE = TEMPLATES / "example_credit.md"


def test_filled_example_passes_the_checker():
    report = check.check_markdown(EXAMPLE)
    assert report.errors == []


def test_blank_template_is_rejected_until_filled():
    report = check.check_markdown(MD_TEMPLATE)
    assert any("placeholder" in e for e in report.errors)


def test_markdown_template_indexes_every_section_and_nothing_else():
    chunks = build_chunks(MD_TEMPLATE.read_text(encoding="utf-8"), level="x", source_file="x.md")
    assert len(chunks) == 14
    assert all(c.keywords for c in chunks)
    joined = " ".join(c.text for c in chunks)
    assert "HOW TO USE" not in joined and "Document Title" not in joined and "Table of Contents" not in joined


def test_word_template_converts_to_the_same_sections(tmp_path):
    md = convert_docx_to_md.convert(DOCX_TEMPLATE)
    word_chunks = build_chunks(md, level="x", source_file="x.md")
    md_chunks = build_chunks(MD_TEMPLATE.read_text(encoding="utf-8"), level="x", source_file="x.md")
    assert [c.section_level2 for c in word_chunks] == [c.section_level2 for c in md_chunks]
    assert all(c.keywords for c in word_chunks)
    assert "Note:" not in md and "HOW TO USE" not in md  # writer guidance removed


def test_word_template_numbers_restart_per_list():
    """Word's List Number style continued one list through the file (steps showed 15, 16, 17)."""
    md = convert_docx_to_md.convert(DOCX_TEMPLATE)
    steps = md[md.index("**Steps:**"):]
    assert steps.splitlines()[2].startswith("1. ")  # first step of task 1
    task2 = md[md.index("## 8."):]
    assert "\n1. " in task2[: task2.index("**Result")]


def test_word_template_keeps_multi_line_notes_whole():
    from docx import Document

    notes = [p.text for p in Document(str(DOCX_TEMPLATE)).paragraphs if p.text.startswith("Note:")]
    assert any(n.endswith("it never grants permission itself.") for n in notes)


def test_converter_handles_every_way_of_numbering(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("Credit", 1)
    doc.add_heading("How to release a hold", 2)
    doc.add_paragraph("Open Customer Orders", style="List Number")  # Word list style
    doc.add_paragraph("Save", style="List Number")
    doc.add_paragraph("Then:")
    doc.add_paragraph("7.\tCheck the customer hold")  # typed number
    doc.add_paragraph("8) Save again")
    doc.add_paragraph("A point", style="List Bullet")
    path = tmp_path / "credit.docx"
    doc.save(path)
    md = convert_docx_to_md.convert(path)
    assert "1. Open Customer Orders\n2. Save" in md
    assert "1. Check the customer hold\n2. Save again" in md  # renumbered after the paragraph
    assert "* A point" in md


def test_excel_template_has_the_columns_the_loader_reads():
    book = openpyxl.load_workbook(XLSX_TEMPLATE)
    assert {"qa_master", "question_variations", "glossary"} <= set(book.sheetnames)
    headers = [c.value for c in book["qa_master"][1]]
    assert headers == check.QA_COLUMNS


def test_excel_template_example_rows_never_go_live(tmp_path):
    copy = tmp_path / "credit.xlsx"
    shutil.copy(XLSX_TEMPLATE, copy)
    records, skipped = _load_one_file(copy, "credit")
    assert records == []  # the example rows are DRAFT
    assert skipped["not approved"] == 2
    assert check.check_excel(copy).errors == []


def test_approved_row_needs_a_real_reviewer(tmp_path):
    copy = tmp_path / "credit.xlsx"
    shutil.copy(XLSX_TEMPLATE, copy)
    book = openpyxl.load_workbook(copy)
    sheet = book["qa_master"]
    headers = [c.value for c in sheet[1]]
    sheet.cell(row=2, column=headers.index("approval_status") + 1, value="APPROVED")
    sheet.cell(row=2, column=headers.index("approved_by") + 1, value="AUTO-DERIVED pending review")
    book.save(copy)
    errors = check.check_excel(copy).errors
    assert any("real reviewer" in e for e in errors)
    records, _ = _load_one_file(copy, "credit")
    assert len(records) == 1  # the loader would use it — which is why the checker blocks it


def test_checker_catches_personal_data_and_leftovers(tmp_path):
    bad = tmp_path / "credit.md"
    bad.write_text(EXAMPLE.read_text(encoding="utf-8").replace(
        "Credit review protects", "Call +91 98765 43210 or mail a.b@company.com. TODO. Credit review protects"),
        encoding="utf-8")
    errors = " ".join(check.check_markdown(bad).errors)
    assert "phone number" in errors and "e-mail" in errors and "TODO" in errors


def test_bad_file_names_are_rejected(tmp_path):
    bad = tmp_path / "Credit Module.md"
    shutil.copy(EXAMPLE, bad)
    assert any("file name" in e for e in check.check_markdown(bad).errors)
